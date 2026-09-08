"""Build unapproved private Final 2-4 baselines from read-only XLSX files."""

from __future__ import annotations

import argparse
import hashlib
import json
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


POINTS_TENTHS = [27] * 12 + [34] * 10 + [42] * 8


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def canonical_digest(value: object) -> str:
    encoded = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def is_number(value: object) -> bool:
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def build_round(source: Path, out_dir: Path, number: int) -> dict[str, object]:
    before_stat = source.stat()
    before = sha256_file(source)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[f"{number}회답안"]
        points_row, header_row = sheet.iter_rows(
            min_row=2, max_row=3, min_col=13, max_col=42, values_only=True
        )
        assert list(header_row) == list(range(1, 31))
        assert [Decimal(str(value)) * 10 for value in points_row] == POINTS_TENTHS

        rows: list[dict[str, object]] = []
        source_rows = sheet.iter_rows(
            min_row=4, max_row=sheet.max_row, min_col=1, max_col=63, values_only=True
        )
        for row_number, source_row in enumerate(source_rows, start=4):
            identity = source_row[1]  # B
            cached_score = source_row[62]  # BK
            if identity in (None, "") or not is_number(cached_score):
                continue

            flags = list(source_row[12:42])  # M:AP
            assert all(flag in (None, "", 0, 1, "0", "1") for flag in flags)
            assert any(flag in (0, 1, "0", "1") for flag in flags), (
                f"final{number} row {row_number}: entirely unmarked row"
            )
            correct = [flag in (1, "1") for flag in flags]
            score_tenths = sum(
                points
                for points, is_correct in zip(POINTS_TENTHS, correct)
                if is_correct
            )
            assert (
                abs(Decimal(str(cached_score)) * 10 - score_tenths)
                < Decimal("0.000001")
            ), f"final{number} row {row_number}: cached source score mismatch"
            rows.append(
                {
                    "id": f"source-row-{row_number}",
                    "ox": "".join("O" if mark else "X" for mark in correct),
                    "score": score_tenths / 10,
                }
            )
    finally:
        workbook.close()

    after = sha256_file(source)
    after_stat = source.stat()
    assert before == after, f"final{number}: source changed during extraction"
    assert (
        before_stat.st_size,
        before_stat.st_mtime_ns,
    ) == (
        after_stat.st_size,
        after_stat.st_mtime_ns,
    ), f"final{number}: source metadata changed during extraction"
    assert rows, f"final{number}: no eligible source rows"

    digest_input = {
        "schemaVersion": 1,
        "exam": f"final{number}",
        "scope": "provided-original-records",
        "rows": rows,
    }
    result = {
        **digest_input,
        "approved": False,
        "version": f"final{number}-{canonical_digest(digest_input)}",
    }
    target = out_dir / f"baseline-final{number}.private.json"
    target.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return {
        "exam": f"final{number}",
        "candidate": target.name,
        "sourceBinding": True,
        "scoreMatch": True,
        "sourceUnchanged": True,
        "approved": False,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for number in (2, 3, 4):
        source = (
            args.source_dir
            / f"2024 황소 초등선발 약점 유형 분석 파이널 {number}회.xlsx"
        )
        results.append(build_round(source, args.out_dir, number))
    print(json.dumps({"pass": True, "results": results}, ensure_ascii=False))


if __name__ == "__main__":
    main()
